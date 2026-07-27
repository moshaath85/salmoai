#!/usr/bin/env python3
"""
Parse the company internal work regulations (لائحة تنظيم العمل) DOCX file
and produce:
  1. company_policy_index.json  – structured article index
  2. Append full text to labor_law_kb.txt
"""

import json
import os
import re
import unicodedata
from docx import Document

DOCX_PATH = "/workspace/uploads/custom_policy_615629.docx"
OUTPUT_JSON = os.path.join(os.path.dirname(__file__), "company_policy_index.json")
KB_PATH = os.path.join(os.path.dirname(__file__), "labor_law_kb.txt")

SYSTEM_NAME = "لائحة تنظيم العمل الداخلية - سالمو"

# Map from Arabic-Indic and Extended Arabic-Indic digits to Western digits
ARABIC_DIGIT_MAP = {
    "٠": "0", "١": "1", "٢": "2", "٣": "3", "٤": "4",
    "٥": "5", "٦": "6", "٧": "7", "٨": "8", "٩": "9",
    "۰": "0", "۱": "1", "۲": "2", "۳": "3", "۴": "4",
    "۵": "5", "۶": "6", "۷": "7", "۸": "8", "۹": "9",
}

# Chapter-to-topic mapping
CHAPTER_TOPIC_MAP = {
    "أحكام عامة": "أحكام عامة",
    "التوظيف": "التوظيف",
    "عقد العمل": "عقود العمل",
    "فترة التجربة": "فترة التجربة",
    "النقل": "النقل والتكليف",
    "الإركاب": "الإركاب والنقل",
    "التدريب والتأهيل": "التدريب والتأهيل",
    "الأجور": "الأجور",
    "تقارير الأداء": "تقارير الأداء",
    "العلاوات": "العلاوات والترقيات",
    "الترقيات": "العلاوات والترقيات",
    "الانتداب": "الانتداب",
    "المزايا والبدلات": "المزايا والبدلات",
    "أيام وساعات العمل": "ساعات العمل",
    "العمل الإضافي": "ساعات العمل",
    "التفتيش الإداري": "التفتيش الإداري",
    "الإجازات": "الإجازات",
    "واجبات المنشأة": "واجبات المنشأة",
    "واجبات العمال": "واجبات العمال",
    "الرعاية الطبية": "الرعاية الطبية",
    "أحكام خاصة بالمرأة": "أحكام خاصة بالمرأة",
    "الخدمات الاجتماعية": "الخدمات الاجتماعية",
    "ضوابط سلوكيات العمل": "ضوابط سلوكيات العمل",
    "انتهاء عقد العمل": "انتهاء عقد العمل",
    "المخالفات والجزاءات": "المخالفات والجزاءات",
    "التظلم": "التظلم",
    "أحكام ختامية": "أحكام ختامية",
}


def normalize_arabic(text: str) -> str:
    """Normalize Arabic presentation forms to standard Arabic letters."""
    # NFKC normalization converts presentation forms (FB50-FDFF, FE70-FEFF) to standard Arabic
    normalized = unicodedata.normalize("NFKC", text)
    return normalized


def arabic_to_western(text: str) -> str:
    """Convert Arabic-Indic digits to Western digits."""
    result = []
    for ch in text:
        result.append(ARABIC_DIGIT_MAP.get(ch, ch))
    return "".join(result)


def extract_article_number(text: str) -> int | None:
    """
    Extract article number from text like 'المادة (١)' or 'اﻟﻤﺎدة (٥)'.
    Only matches when the article header is at the START of the paragraph
    (possibly preceded by whitespace or a period/number).
    Returns the article number as int, or None if not an article header.
    """
    norm = normalize_arabic(text.strip())
    # Match article header at the beginning of the text (the paragraph is primarily an article header)
    # Allow optional leading characters like period, number, whitespace
    pattern = r"^[\s.\d٠-٩۰-۹]*المادة\s*\(([٠-٩۰-۹0-9]+)\)"
    m = re.match(pattern, norm)
    if m:
        num_str = arabic_to_western(m.group(1))
        try:
            return int(num_str)
        except ValueError:
            return None
    return None


def normalize_chapter_name(raw: str) -> str:
    """Normalize a chapter heading to match our topic map keys."""
    norm = normalize_arabic(raw.strip())
    # Remove extra whitespace
    norm = re.sub(r"\s+", " ", norm).strip()
    return norm


def find_topic_for_chapter(chapter_name: str) -> str:
    """Find the best matching topic for a chapter name."""
    norm = normalize_chapter_name(chapter_name)
    # Direct match
    if norm in CHAPTER_TOPIC_MAP:
        return CHAPTER_TOPIC_MAP[norm]
    # Fuzzy match: check if any key is contained in the chapter name
    for key, topic in CHAPTER_TOPIC_MAP.items():
        if key in norm or norm in key:
            return topic
    return norm  # fallback to chapter name itself


def parse_penalty_tables(doc: Document) -> str:
    """Extract penalty tables as formatted text."""
    lines = []
    for ti, table in enumerate(doc.tables):
        rows_data = []
        for row in table.rows:
            cells = [normalize_arabic(c.text.strip()) for c in row.cells]
            rows_data.append(cells)

        if not rows_data:
            continue

        # Build a readable text representation
        if ti == 0:
            lines.append("\n\nجداول المخالفات والجزاءات\n")
            lines.append("أولاً: مخالفات تتعلق بمواعيد العمل:\n")
        elif ti == 3:
            lines.append("\nثانياً: مخالفات تتعلق بتنظيم العمل:\n")
        elif ti == 6:
            lines.append("\nثالثاً: مخالفات تتعلق بسلوك العامل:\n")

        for ri, row_cells in enumerate(rows_data):
            if ri < 2:  # Skip header rows (already described)
                continue
            # Format: violation number | violation type | 1st | 2nd | 3rd | 4th
            # Cells are in reverse order (RTL): [4th, 3rd, 2nd, 1st, violation, number]
            if len(row_cells) >= 6:
                num = row_cells[5]
                violation = row_cells[4]
                first = row_cells[3]
                second = row_cells[2]
                third = row_cells[1]
                fourth = row_cells[0]
                lines.append(
                    f"  {num}. {violation} | المرة الأولى: {first} | المرة الثانية: {second} | المرة الثالثة: {third} | المرة الرابعة: {fourth}"
                )

    return "\n".join(lines)


def main():
    doc = Document(DOCX_PATH)

    articles = []
    current_chapter = "مقدمة"
    current_article_num = None
    current_article_chapter = None  # Chapter at the time the article started
    current_article_text_parts = []
    introduction_text_parts = []

    # Track where articles end (before penalty tables section)
    # Articles are in paragraphs, penalty tables are in doc.tables

    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue

        norm_text = normalize_arabic(text)

        # Check if this is a chapter heading
        if para.style.name == "Heading 1":
            chapter_name = normalize_chapter_name(text)
            current_chapter = chapter_name
            continue

        # Check if this is an article header
        art_num = extract_article_number(text)
        if art_num is not None:
            # Save previous article if any
            if current_article_num is not None:
                full_text = "\n".join(current_article_text_parts).strip()
                ch = current_article_chapter or current_chapter
                topic = find_topic_for_chapter(ch)
                articles.append({
                    "article_number": current_article_num,
                    "article_name": f"المادة ({current_article_num})",
                    "text": full_text,
                    "chapter": ch,
                    "section": "",
                    "cancelled": False,
                    "system": SYSTEM_NAME,
                    "topic": topic,
                })

            current_article_num = art_num
            current_article_chapter = current_chapter  # Record chapter when article starts
            current_article_text_parts = []

            # Check if the paragraph contains more than just the article header
            # e.g., "المادة (٣٠)\nsome text"
            remainder = re.sub(r"المادة\s*\([٠-٩۰-۹0-9]+\)\s*", "", norm_text).strip()
            if remainder:
                current_article_text_parts.append(remainder)
            continue

        # Check if text contains an inline article header (e.g., "... اﻟﻤﺎدة (۳۰)")
        # Some paragraphs end with an article header for the next article
        inline_match = re.search(r"المادة\s*\(([٠-٩۰-۹0-9]+)\)\s*$", norm_text)
        if inline_match and current_article_num is not None:
            # Split: text before is part of current article, the header starts a new one
            before = norm_text[:inline_match.start()].strip()
            if before:
                current_article_text_parts.append(before)

            # Save current article
            full_text = "\n".join(current_article_text_parts).strip()
            ch = current_article_chapter or current_chapter
            topic = find_topic_for_chapter(ch)
            articles.append({
                "article_number": current_article_num,
                "article_name": f"المادة ({current_article_num})",
                "text": full_text,
                "chapter": ch,
                "section": "",
                "cancelled": False,
                "system": SYSTEM_NAME,
                "topic": topic,
            })

            new_num_str = arabic_to_western(inline_match.group(1))
            current_article_num = int(new_num_str)
            current_article_chapter = current_chapter  # Record chapter when new article starts
            current_article_text_parts = []
            continue

        # If we're inside an article, accumulate text
        if current_article_num is not None:
            current_article_text_parts.append(norm_text)
        elif i < 20:
            # Before first article - introduction
            introduction_text_parts.append(norm_text)

    # Save the last article
    if current_article_num is not None:
        full_text = "\n".join(current_article_text_parts).strip()
        ch = current_article_chapter or current_chapter
        topic = find_topic_for_chapter(ch)
        articles.append({
            "article_number": current_article_num,
            "article_name": f"المادة ({current_article_num})",
            "text": full_text,
            "chapter": ch,
            "section": "",
            "cancelled": False,
            "system": SYSTEM_NAME,
            "topic": topic,
        })

    # Parse penalty tables and add as a special article
    penalty_text = parse_penalty_tables(doc)
    if penalty_text.strip():
        articles.append({
            "article_number": 0,
            "article_name": "جداول المخالفات والجزاءات",
            "text": penalty_text.strip(),
            "chapter": "المخالفات والجزاءات",
            "section": "جداول المخالفات والجزاءات",
            "cancelled": False,
            "system": SYSTEM_NAME,
            "topic": "المخالفات والجزاءات",
        })

    # Write JSON output
    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(articles, f, ensure_ascii=False, indent=2)

    # Build full text for KB
    kb_lines = []
    kb_lines.append("\n\n" + "=" * 80)
    kb_lines.append(f"لائحة تنظيم العمل الداخلية - مؤسسة سالمو سوليوشنز")
    kb_lines.append("=" * 80)
    kb_lines.append("")

    # Introduction
    if introduction_text_parts:
        kb_lines.append("مقدمة:")
        kb_lines.extend(introduction_text_parts)
        kb_lines.append("")

    # Articles
    current_ch = None
    for art in articles:
        if art["article_number"] == 0:
            # Penalty tables
            kb_lines.append("")
            kb_lines.append(art["text"])
            continue

        if art["chapter"] != current_ch:
            current_ch = art["chapter"]
            kb_lines.append("")
            kb_lines.append(f"--- {current_ch} ---")
            kb_lines.append("")

        kb_lines.append(f"{art['article_name']}:")
        kb_lines.append(art["text"])
        kb_lines.append("")

    full_kb_text = "\n".join(kb_lines)

    # Append to KB file
    with open(KB_PATH, "a", encoding="utf-8") as f:
        f.write(full_kb_text)

    # Print summary
    chapters = sorted(set(a["chapter"] for a in articles if a["article_number"] > 0))
    print(f"✅ Successfully parsed company policy document")
    print(f"   Articles extracted: {len([a for a in articles if a['article_number'] > 0])}")
    print(f"   Penalty tables: {'Yes' if penalty_text.strip() else 'No'}")
    print(f"   Total entries in JSON: {len(articles)}")
    print(f"   Chapters found ({len(chapters)}):")
    for ch in chapters:
        count = len([a for a in articles if a["chapter"] == ch and a["article_number"] > 0])
        print(f"     - {ch} ({count} articles)")
    print(f"   Output JSON: {OUTPUT_JSON}")
    print(f"   KB appended: {KB_PATH}")
    print(f"   KB text added: {len(full_kb_text)} chars")


if __name__ == "__main__":
    main()